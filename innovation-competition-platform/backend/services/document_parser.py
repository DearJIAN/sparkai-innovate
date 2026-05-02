import os
import logging

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    '.txt': 'parse_txt',
    '.pdf': 'parse_pdf',
    '.docx': 'parse_docx',
    '.pptx': 'parse_pptx',
}

UNSUPPORTED_EXTENSIONS = {
    '.doc': '旧版 Word(.doc)格式暂不支持，请转换为 .docx 后重新上传',
    '.ppt': '旧版 PPT(.ppt)格式暂不支持，请转换为 .pptx 后重新上传',
    '.jpg': '图片文件暂不支持文本提取，后续可接入 OCR',
    '.jpeg': '图片文件暂不支持文本提取，后续可接入 OCR',
    '.png': '图片文件暂不支持文本提取，后续可接入 OCR',
    '.gif': '图片文件暂不支持文本提取，后续可接入 OCR',
    '.bmp': '图片文件暂不支持文本提取，后续可接入 OCR',
    '.mp4': '视频文件暂不支持文本提取，后续可接入 ASR',
    '.avi': '视频文件暂不支持文本提取，后续可接入 ASR',
    '.mov': '视频文件暂不支持文本提取，后续可接入 ASR',
    '.wmv': '视频文件暂不支持文本提取，后续可接入 ASR',
    '.mkv': '视频文件暂不支持文本提取，后续可接入 ASR',
    '.zip': '压缩包请先解压后上传内部文档',
    '.rar': '压缩包请先解压后上传内部文档',
    '.7z': '压缩包请先解压后上传内部文档',
}


def parse_txt(file_path):
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        if not content.strip():
            return None, '文件内容为空'
        return content, None
    except UnicodeDecodeError:
        try:
            with open(file_path, 'r', encoding='gbk') as f:
                content = f.read()
            if not content.strip():
                return None, '文件内容为空'
            return content, None
        except Exception as e:
            return None, f'文本文件读取失败: {str(e)}'
    except Exception as e:
        return None, f'文本文件读取失败: {str(e)}'


def parse_pdf(file_path):
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        texts = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and text.strip():
                texts.append(text.strip())
        if not texts:
            return None, 'PDF 文件中未提取到文本内容（可能是扫描件或纯图片 PDF）'
        return '\n\n'.join(texts), None
    except ImportError:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            texts = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    texts.append(text.strip())
            if not texts:
                return None, 'PDF 文件中未提取到文本内容（可能是扫描件或纯图片 PDF）'
            return '\n\n'.join(texts), None
        except ImportError:
            return None, 'PDF 解析库未安装，请安装 pypdf 或 PyPDF2'
        except Exception as e:
            return None, f'PDF 文件解析失败: {str(e)}'
    except Exception as e:
        return None, f'PDF 文件解析失败: {str(e)}'


def parse_docx(file_path):
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(' | '.join(row_texts))
        if not paragraphs:
            return None, 'DOCX 文件中未提取到文本内容'
        return '\n\n'.join(paragraphs), None
    except ImportError:
        return None, 'DOCX 解析库未安装，请安装 python-docx'
    except Exception as e:
        return None, f'DOCX 文件解析失败: {str(e)}'


def parse_pptx(file_path):
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        slides_text = []
        for i, slide in enumerate(prs.slides):
            slide_parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_parts.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_texts:
                            slide_parts.append(' | '.join(row_texts))
            if slide_parts:
                slides_text.append(f'[幻灯片 {i + 1}]\n' + '\n'.join(slide_parts))
        if not slides_text:
            return None, 'PPTX 文件中未提取到文本内容'
        return '\n\n'.join(slides_text), None
    except ImportError:
        return None, 'PPTX 解析库未安装，请安装 python-pptx'
    except Exception as e:
        return None, f'PPTX 文件解析失败: {str(e)}'


def parse_file(file_path, original_name=None):
    ext = os.path.splitext(file_path)[1].lower()
    if original_name:
        ext_from_name = os.path.splitext(original_name)[1].lower()
        ext = ext_from_name if ext_from_name in SUPPORTED_EXTENSIONS or ext_from_name in UNSUPPORTED_EXTENSIONS else ext

    if ext in UNSUPPORTED_EXTENSIONS:
        return None, UNSUPPORTED_EXTENSIONS[ext], 'unsupported'

    if ext not in SUPPORTED_EXTENSIONS:
        return None, f'不支持的文件格式: {ext}，当前支持 .txt/.pdf/.docx/.pptx', 'unsupported'

    parser_name = SUPPORTED_EXTENSIONS[ext]
    parser_func = globals().get(parser_name)
    if not parser_func:
        return None, f'解析器 {parser_name} 未找到', 'error'

    content, error = parser_func(file_path)
    if error:
        return None, error, 'error'

    return content, None, 'success'


def parse_project_files(files, base_upload_dir):
    results = []
    for f in files:
        file_path = os.path.join(base_upload_dir, f.file_path) if not os.path.isabs(f.file_path) else f.file_path
        if not os.path.exists(file_path):
            results.append({
                'file_id': f.id,
                'file_name': f.original_name,
                'status': 'error',
                'reason': '文件在服务器上不存在',
                'content': None
            })
            continue

        content, error, status = parse_file(file_path, f.original_name)
        results.append({
            'file_id': f.id,
            'file_name': f.original_name,
            'material_type': getattr(f, 'material_type', None),
            'status': status,
            'reason': error,
            'content': content
        })

    return results


def parse_registration_materials(materials, base_upload_dir):
    results = []
    for m in materials:
        file_path = os.path.join(base_upload_dir, m.file_path) if not os.path.isabs(m.file_path) else m.file_path
        if not os.path.exists(file_path):
            results.append({
                'file_id': m.id,
                'file_name': m.original_name,
                'status': 'error',
                'reason': '文件在服务器上不存在',
                'content': None
            })
            continue

        content, error, status = parse_file(file_path, m.original_name)
        results.append({
            'file_id': m.id,
            'file_name': m.original_name,
            'material_type': getattr(m, 'material_type', None),
            'status': status,
            'reason': error,
            'content': content
        })

    return results


def get_supported_extensions():
    return list(SUPPORTED_EXTENSIONS.keys())


def get_unsupported_extensions():
    return list(UNSUPPORTED_EXTENSIONS.keys())
